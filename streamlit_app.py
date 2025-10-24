import streamlit as st
import pandas as pd
import openpyxl
from openpyxl import load_workbook
import re
from datetime import datetime
import io

def extract_url_from_hyperlink_formula(formula_text):
    """
    Izvelk URL no HYPERLINK formulas
    Atpazīst: =HYPERLINK("URL", "text") 
    """
    if not formula_text or not isinstance(formula_text, str):
        return None
        
    # Pattern to match =HYPERLINK("URL", "text")
    pattern = r'=HYPERLINK\("([^"]+)"'
    match = re.search(pattern, formula_text)
    
    if match:
        return match.group(1)
    return None

def read_excel_with_hyperlinks(file_path):
    """
    Lasa Excel failu un ekstraktē URL no HYPERLINK formulām
    Ekstraktē gan YouTube linkus (Links kolonna), gan MP3 linkus (Dwnld. kolonna)
    """
    # Read with pandas for basic data
    df = pd.read_excel(file_path)
    
    # Read with openpyxl to get formulas
    wb = load_workbook(file_path, data_only=False)  # data_only=False to get formulas
    ws = wb.active
    
    # Find Links and Dwnld columns
    links_col = None
    dwnld_col = None
    
    for col in range(1, ws.max_column + 1):
        cell_value = ws.cell(row=1, column=col).value
        if cell_value == 'Links':
            links_col = col
        elif cell_value == 'Dwnld.':
            dwnld_col = col
    
    # Extract YouTube URLs from Links column
    if links_col and 'Links' in df.columns:
        extracted_youtube_urls = []
        
        for row in range(2, ws.max_row + 1):  # Skip header
            cell = ws.cell(row=row, column=links_col)
            cell_value = cell.value
            
            # Try to extract URL from formula
            url = extract_url_from_hyperlink_formula(str(cell_value))
            
            # If no URL extracted, try Direct URL column as fallback
            if not url and 'Direct URL' in df.columns:
                try:
                    direct_url = df.loc[row-2, 'Direct URL']  # row-2 because pandas is 0-indexed
                    if pd.notna(direct_url):
                        url = direct_url
                except:
                    pass
            
            extracted_youtube_urls.append(url)
        
        # Add extracted YouTube URLs to dataframe
        df['YouTube_URL'] = extracted_youtube_urls[:len(df)]
    
    # Extract MP3 URLs from Dwnld. column
    if dwnld_col and 'Dwnld.' in df.columns:
        extracted_mp3_urls = []
        
        for row in range(2, ws.max_row + 1):  # Skip header
            cell = ws.cell(row=row, column=dwnld_col)
            cell_value = cell.value
            
            # Try to extract URL from formula
            mp3_url = extract_url_from_hyperlink_formula(str(cell_value))
            extracted_mp3_urls.append(mp3_url)
        
        # Add extracted MP3 URLs to dataframe
        df['MP3_URL'] = extracted_mp3_urls[:len(df)]
    
    return df

def create_clickable_link(url, text="🔗 Atvērt"):
    """Izveido klikšķināmu linku Streamlit"""
    if url and isinstance(url, str) and url.startswith('http'):
        return f'<a href="{url}" target="_blank">{text}</a>'
    return "Nav pieejams"

def search_mp3_in_drive_by_name(file_name, drive_folder_id="16qboBQJSIHhto5TioupnkfsE6DrQpLV7"):
    """
    Meklē MP3 failu Google Drive pēc nosaukuma un ģenerē linku
    (Šis ir piemērs - reālajā izmantošanā vajag Google Drive API atslēgas)
    """
    if not file_name:
        return None
    
    # Remove file extension and clean the name
    clean_name = str(file_name).replace('.mp3', '').replace('.wav', '').strip()
    
    # For demo, generate potential Google Drive link format
    # In real implementation, this would use Google Drive API
    if clean_name:
        # This is a placeholder - real implementation would search Drive API
        return f"https://drive.google.com/drive/folders/{drive_folder_id}?q={clean_name.replace(' ', '%20')}"
    
    return None

def generate_mp3_search_link(original_filename):
    """Ģenerē meklēšanas linku MP3 failam"""
    if not original_filename:
        return None
    
    # Clean filename for search
    clean_name = str(original_filename)
    # Remove common patterns
    for pattern in [' - \\d{4}-\\d{2}-\\d{2}', '\\s*-\\s*Sri Prem Prayojan.*', '\\.mp3$', '\\.wav$']:
        clean_name = re.sub(pattern, '', clean_name)
    
    clean_name = clean_name.strip()
    
    if clean_name:
        # Generate Google Drive search URL
        encoded_name = clean_name.replace(' ', '%20')
        return f"https://drive.google.com/drive/search?q={encoded_name}"
    
    return None

def main():
    st.set_page_config(
        page_title="CA Datu Bāze - Hipersaišu Lasītājs",
        page_icon="🎵",
        layout="wide"
    )
    
    st.title("🎵 Chaitanya Academy Datu Bāze")
    st.subheader("📁 Audio failu katalogs ar hipersaišu atbalstu")
    
    # File uploader
    uploaded_file = st.file_uploader(
        "Augšupielādējiet Excel failu", 
        type=['xlsx', 'xls'],
        help="Augšupielādējiet CA datu bāzes Excel failu"
    )
    
    if uploaded_file is not None:
        try:
            with st.spinner('Apstrādā failu un ekstraktē hipersaites...'):
                # Save uploaded file temporarily
                with open('temp_file.xlsx', 'wb') as f:
                    f.write(uploaded_file.getvalue())
                
                # Read with hyperlink extraction
                df = read_excel_with_hyperlinks('temp_file.xlsx')
            
            st.success(f"✅ Fails veiksmīgi augšupielādēts! Atrasti {len(df)} ieraksti")
            
            # Show extraction statistics
            youtube_count = 0
            mp3_count = 0
            
            if 'YouTube_URL' in df.columns:
                youtube_count = df['YouTube_URL'].notna().sum()
                
            if 'MP3_URL' in df.columns:
                mp3_count = df['MP3_URL'].notna().sum()
            
            if youtube_count > 0 or mp3_count > 0:
                st.info(f"🔗 Ekstraktēti: {youtube_count} YouTube linki, {mp3_count} MP3 linki no {len(df)} ierakstiem")
            
            # Sidebar filters
            st.sidebar.header("🔍 Filtri")
            
            # Source filter
            if 'Source' in df.columns:
                sources = ['Visi'] + sorted(df['Source'].dropna().unique().tolist())
                selected_source = st.sidebar.selectbox("Avots:", sources)
                
                if selected_source != 'Visi':
                    df = df[df['Source'] == selected_source]
            
            # Subject filter
            if 'Subject' in df.columns:
                subjects = ['Visi'] + sorted(df['Subject'].dropna().unique().tolist())
                selected_subject = st.sidebar.selectbox("Temats:", subjects)
                
                if selected_subject != 'Visi':
                    df = df[df['Subject'] == selected_subject]
            
            # Search
            search_term = st.sidebar.text_input("🔍 Meklēt:")
            if search_term:
                mask = df.astype(str).apply(lambda x: x.str.contains(search_term, case=False, na=False)).any(axis=1)
                df = df[mask]
            
            st.write(f"**Rāda {len(df)} ierakstus**")
            
            # Display data
            if len(df) > 0:
                # Prepare display dataframe
                display_df = df.copy()
                
                # Create clickable links
                if 'YouTube_URL' in display_df.columns:
                    display_df['YouTube saite'] = display_df['YouTube_URL'].apply(
                        lambda x: create_clickable_link(x, "▶️ YouTube") if x else "❌ Nav"
                    )
                
                if 'MP3_URL' in display_df.columns:
                    display_df['MP3 fails'] = display_df['MP3_URL'].apply(
                        lambda x: create_clickable_link(x, "🎵 MP3") if x else "❌ Nav"
                    )
                
                # Select important columns for display
                important_cols = []
                for col in ['Original file name', 'Subject', 'Date', 'Source', 'Length', 'YouTube saite', 'MP3 fails']:
                    if col in display_df.columns:
                        important_cols.append(col)
                
                if important_cols:
                    # Display with clickable links
                    st.markdown(
                        display_df[important_cols].to_html(escape=False, index=False), 
                        unsafe_allow_html=True
                    )
                else:
                    st.dataframe(display_df)
                
                # Download options
                st.subheader("📥 Lejupielāde")
                
                col1, col2 = st.columns(2)
                
                with col1:
                    # CSV download
                    csv = df.to_csv(index=False, encoding='utf-8')
                    st.download_button(
                        label="💾 Lejupielādēt CSV",
                        data=csv,
                        file_name=f"CA_database_{datetime.now().strftime('%Y%m%d_%H%M')}.csv",
                        mime="text/csv"
                    )
                
                with col2:
                    # Excel download
                    output = io.BytesIO()
                    with pd.ExcelWriter(output, engine='openpyxl') as writer:
                        df.to_excel(writer, index=False, sheet_name='CA_Data')
                    
                    st.download_button(
                        label="📊 Lejupielādēt Excel",
                        data=output.getvalue(),
                        file_name=f"CA_database_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
            
            else:
                st.warning("⚠️ Nav atrasti ieraksti pēc jūsu filtriem")
                
        except Exception as e:
            st.error(f"❌ Kļūda faila apstrādē: {str(e)}")
            st.info("Pārbaudiet, vai fails ir pareizā Excel formātā")
    
    else:
        st.info("👆 Lūdzu augšupielādējiet Excel failu, lai sāktu")
        
        # Instructions
        with st.expander("📖 Instrukcijas"):
            st.markdown("""
            ### Kā izmantot šo aplikāciju:
            
            1. **Augšupielādējiet Excel failu** ar CA datu bāzi
            2. **Izvēlieties filtrus** sānjoslā (avots, temats)
            3. **Meklējiet** konkrētu saturu
            4. **Klikšķiniet** uz ▶️ YouTube pogām, lai atvērtu video
            5. **Klikšķiniet** uz 🎵 MP3 pogām, lai atvērtu audio failus
            6. **Lejupielādējiet** rezultātus CSV vai Excel formātā
            
            ### Funkcijas:
            - ✅ **Hipersaišu ekstraktēšana** no Excel HYPERLINK formulām
            - 🔍 **Meklēšana un filtrēšana** pēc dažādiem kritērijiem  
            - 🔗 **Klikšķināmi YouTube linki** uz video saturu
            - 🎵 **Tiešie MP3 linki** uz Google Drive audio failiem
            - 📥 **Eksports** CSV un Excel formātos
            - 📱 **Atsaucīgs dizains** darbam dažādās ierīcēs
            """)

if __name__ == "__main__":
    main()
